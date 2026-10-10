"""Excel export of the final database for readers who open data in a spreadsheet.

    python -m investordb.cli export-xlsx     -> data/processed/investori_cz_sk.xlsx

Sheets: Investori (one row per investor, readable labels), Tvrdenia (every verified claim behind the values, with a
clickable source and the verbatim quote), Popis stĺpcov (data dictionary), O súbore (what this is, how it was made).
Plain data, no formulas: every number comes from investors_refined.csv.
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from investordb.datadict import CLAIM_COLUMNS, CODES, INVESTOR_COLUMNS, SECTOR_SK, STAGE_SK, TYPE_SK
from investordb.evidence import domain
from investordb.refine import PROCESSED, _value, rebuild_all

OUT = PROCESSED / "investori_cz_sk.xlsx"
REPO = "https://github.com/Matush77/Clientelo-task-a"
EXPLORER = "https://claude.ai/artifact/QNrw3gMzFoMqE9kSpoD1SS"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1A44A3")
BAND_FILL = PatternFill("solid", fgColor="F2F5F9")
THIN = Side(style="thin", color="C9D2DD")
NUMBER_COLS = {"ticket_min_eur", "ticket_max_eur", "total_capital_eur", "n_investments", "n_investments_36m",
               "capital_approx"}
LINK_COLS = {"last_investment_source", "registry_url", "website"}
ORIGIN = {"refined": "spresnenie (D38)", "gapfill": "doplnenie (D41)"}
FIELD_SK = {"investments": "investícia", "funds": "fond", "total_capital": "AUM", "ticket": "tiket", "sectors": "sektory",
            "stages": "štádiá", "hq_country": "sídlo", "investor_type": "typ", "identity": "identita"}
FIELD_ORDER = list(FIELD_SK)


def _labels(codes: str, table: dict) -> str:
    return ", ".join(table.get(c, c) for c in codes.split(",") if c)


def _cell_value(col: str, row: dict):
    v = row.get(col, "")
    if col in NUMBER_COLS:
        return float(v) if v not in ("", None) else None
    if col == "sectors":
        return _labels(v, SECTOR_SK)
    if col == "stages":
        return _labels(v, STAGE_SK)
    if col == "investor_types":
        return _labels(v, TYPE_SK)
    if col in ("sectors_basis", "stages_basis"):
        return {"stated": "uvádza investor", "inferred": "odvodené z portfólia"}.get(v, v)
    if col == "not_public":
        return ", ".join({"sectors": "sektory", "stages": "štádiá", "ticket": "tiket",
                          "total_capital": "celkový kapitál"}.get(f, f) for f in v.split("; ") if f)
    return v


def _claim_text(c: dict) -> str:
    v = _value(c) if c.get("value", "").startswith("{") else None
    if c["field"] == "investments" and v:
        parts = [v.get("company"), c.get("event_date", "")[:10] if c.get("deal_context") == "deal" else "bez dátumu obchodu",
                 v.get("amount")]
        return " · ".join(str(p) for p in parts if p)
    if c["field"] == "funds" and v:
        status = {"final_close": "uzavretý", "first_close": "prvé uzavretie", "target": "cieľ"}.get(v.get("status"), "")
        return f"{v.get('name') or 'fond'}: {v.get('size') or '?'}" + (f" ({status})" if status else "")
    if c["field"] == "ticket" and v:
        return " – ".join(str(x) for x in (v.get("min"), v.get("max")) if x)
    if c["field"] in ("sectors", "stages"):
        codes = json.loads(c["value"]) if c.get("value") else []
        codes = [codes] if isinstance(codes, str) else codes or []
        table = SECTOR_SK if c["field"] == "sectors" else STAGE_SK
        return ", ".join(table.get(x, x) for x in codes)
    return c.get("value_text") or (json.dumps(v, ensure_ascii=False) if v else c.get("value", ""))


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def _finish(ws, n_cols: int, wrap_cols: set[int] = frozenset()) -> None:
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=n_cols), 2):
        for cell in row:
            if not cell.hyperlink:  # links keep their own blue, underlined font
                cell.font = Font(name=FONT, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
            cell.border = Border(bottom=THIN)
            if r % 2 == 0:
                cell.fill = BAND_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{ws.max_row}"


def _link(cell, url: str, label: str | None = None) -> None:
    if url and url.startswith("http"):
        cell.value = label or url
        cell.hyperlink = url
        cell.font = Font(name=FONT, color="1A44A3", underline="single")


def build(as_of: date) -> Path:
    result = rebuild_all(as_of)
    rows = sorted(result["rows"].values(), key=lambda r: r["name"].lower())
    wb = Workbook()

    # --- investors
    ws = wb.active
    ws.title = "Investori"
    cols = INVESTOR_COLUMNS
    widths = {"name": 26, "funds": 46, "funds_target": 34, "capital_note": 34, "refine_notes": 46, "explanation": 34,
              "legal_name": 30, "last_investment_source": 22, "registry_url": 16, "website": 26, "sectors": 30}
    _header(ws, [h for _, h, _ in cols], [widths.get(c, 16) for c, _, _ in cols])
    for row in rows:
        ws.append([_cell_value(c, row) for c, _, _ in cols])
        r = ws.max_row
        for i, (c, _, _) in enumerate(cols, 1):
            cell = ws.cell(row=r, column=i)
            if c in LINK_COLS:
                _link(cell, row.get(c, ""), domain(row.get(c, "")) if c != "registry_url" else "register")
            elif c in ("ticket_min_eur", "ticket_max_eur", "total_capital_eur"):
                cell.number_format = '#,##0 "€"'
    wrap = {i for i, (c, _, _) in enumerate(cols, 1) if c in ("funds", "funds_target", "capital_note", "refine_notes",
                                                              "explanation", "sectors", "legal_name")}
    _finish(ws, len(cols), wrap)

    # --- claims behind the values
    wc = wb.create_sheet("Tvrdenia")
    heads = ["Investor", "Pole", "Hodnota", "Doslovná citácia (pôvodný jazyk)", "Zdroj", "Publikované", "Typ zdroja",
             "Odvodenie", "Kontext", "Pôvod tvrdenia", "Zhoda citácie"]
    _header(wc, heads, [26, 12, 40, 70, 26, 12, 14, 12, 14, 18, 10])
    tier = {"T1": "register", "T2": "web investora", "T3": "tlač"}
    context = {"deal": "správa o obchode", "mention": "zmienka", "exit": "exit"}
    for row in rows:
        claims = [c for c in result["merged"][row["candidate_id"]] if c["auto_check"] == "ok" and c["field"] in FIELD_SK
                  and not (c["field"] == "investments" and c.get("attributed") == "0")]
        claims.sort(key=lambda c: (FIELD_ORDER.index(c["field"]), -int((c.get("event_date") or "0000")[:4] or 0)))
        for c in claims:
            origin = c.get("origin", "")
            wc.append([row["name"], FIELD_SK[c["field"]], _claim_text(c), c["quote"], None, c.get("published_date") or "",
                       tier.get(c.get("source_tier"), c.get("source_tier")),
                       {"stated": "uvádza zdroj", "inferred": "odvodené"}.get(c.get("derivation"), c.get("derivation")),
                       context.get(c.get("deal_context"), ""), ORIGIN.get(origin, "zber dôkazov"),
                       float(c["quote_score"]) if c.get("quote_score") else None])
            _link(wc.cell(row=wc.max_row, column=5), c["source_url"], domain(c["source_url"]))
    _finish(wc, len(heads), {3, 4})

    # --- data dictionary
    wd = wb.create_sheet("Popis stĺpcov")
    _header(wd, ["Hárok / skupina", "Stĺpec alebo kód", "Hlavička v Exceli", "Popis"], [22, 24, 28, 90])
    for c, h, d in INVESTOR_COLUMNS:
        wd.append(["Investori", c, h, d.replace("`", "")])
    for c, d in CLAIM_COLUMNS:
        wd.append(["claims*.csv", c, "", d.replace("`", "")])
    for group, codes in CODES.items():
        for c, d in codes:
            wd.append([group, c, "", d])
    _finish(wd, 4, {4})
    wd.freeze_panes = "A2"

    # --- about
    wa = wb.create_sheet("O súbore")
    about = [
        ("Databáza VC investorov so sídlom v Česku a na Slovensku", ""),
        ("Stav k", result["rows"][rows[0]["candidate_id"]]["as_of"]),
        ("Zaradených investorov", str(len(rows))),
        ("Čo je v súbore", "Výsledná databáza po spresnení (D38) a doplnení polí (D41). Každá hodnota má v hárku "
                           "Tvrdenia zdroj a doslovnú citáciu, ktorú program overil na stránke zdroja."),
        ("Prázdne bunky", "Údaj sa nenašiel. Ak je pole v stĺpci „Verejne neuvedené“, agent ho hľadal a verejne sa "
                          "neuvádza – nič nie je odhadnuté."),
        ("Sumy v EUR", "Prepočet kurzom ECB z posledného dňa pred dátumom „Stav k“."),
        ("Repozitár", REPO),
        ("Prehliadač s citáciami", EXPLORER),
        ("Vygenerované", "python -m investordb.cli export-xlsx"),
    ]
    wa.column_dimensions["A"].width = 26
    wa.column_dimensions["B"].width = 100
    for k, v in about:
        wa.append([k, v])
        cell = wa.cell(row=wa.max_row, column=2)
        if v.startswith("http"):
            _link(cell, v)
    for row in wa.iter_rows():
        for cell in row:
            if not cell.hyperlink:
                cell.font = Font(name=FONT, bold=cell.column == 1 or cell.row == 1, size=14 if cell.row == 1 else 11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.move_sheet("O súbore", offset=-3)
    wb.active = 1
    wb.save(OUT)
    return OUT
